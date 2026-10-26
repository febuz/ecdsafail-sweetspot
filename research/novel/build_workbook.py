#!/usr/bin/env python3
"""Build the technique-analysis workbook (docs/technieken_analyse.xlsx).

Sheets: Legenda, Gebruikt (all solutions), Ongebruikt, Tekstboek (2-qubit gates to 256-bit blocks), Weging
(extra properties + weight model), Gelijkheid (Gower similarity to the used set), Ontwerp (novel design, formulas),
Vergelijking (leader, Google, published, novel), Mijlpalen.
Inputs: scratchpad used_techniques.json / milestones.json (mined from 154 submission notes), unused_techniques.py,
textbook_techniques.py. Every judgement value is marked as such in Legenda.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from textbook_techniques import TEXTBOOK
from unused_techniques import UNUSED

SCR = sys.argv[1]
OUT = sys.argv[2]
F = "Arial"
HDR = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
BOLD = Font(name=F, bold=True)
GREEN = Font(name=F, color="008000")
YELLOW = PatternFill("solid", fgColor="FFFF00")

COLS = ["ID", "Techniek", "Categorie", "Vakgebied", "Jaar", "Bron", "Primair doel", "Vast aantal qubits nodig",
        "Extra qubits @256 (schatting)", "Ancilla-type", "Coherentie-eis (1-5)", "Verstrengeling met data nodig",
        "Mid-circuit meting", "Klassieke feedforward", "Exact / benaderd", "Foutkans", "Toffoli-schaal", "Qubit-schaal",
        "Diepte-schaal", "Challenge-gateset ok", "Rotaties nodig", "Data-afhankelijk pad", "Uncompute-methode",
        "Vooraf berekende tabellen", "Fase-risico", "Complexiteit (1-5)", "Volwassenheid (1-5)",
        "T-effect op leader (%)", "Q-effect op leader (qubits)", "In leader", "Elders gebruikt", "Opmerking"]
# numeric feature block (Gower): column index in COLS (1-based) -> formula template on that cell
FEATS = [
    ("vast", 8, 'IF({c}="ja",1,0)'),
    ("ancilla", 10, 'IF({c}="schoon",1,IF({c}="vuil",0.5,0))'),
    ("coherentie", 11, "{c}/5"),
    ("verstrengeling", 12, 'IF({c}="ja",1,0)'),
    ("meting", 13, 'IF({c}="ja",1,0)'),
    ("feedforward", 14, 'IF({c}="ja",1,0)'),
    ("exact", 15, 'IF({c}="exact",1,0)'),
    ("gateset", 20, 'IF({c}="ja",1,0)'),
    ("rotaties", 21, 'IF({c}="ja",1,0)'),
    ("data-afh", 22, 'IF({c}="ja",1,0)'),
    ("uncompute", 23, 'IF({c}="Bennett",1,IF({c}="meting",0.5,0))'),
    ("tabellen", 24, 'IF({c}="ja",1,0)'),
    ("complexiteit", 26, "{c}/5"),
    ("volwassenheid", 27, "{c}/5"),
]
NF = len(FEATS)
FEAT0 = len(COLS) + 2  # first helper column (one blank column gap)

CAT_DEFAULTS = {  # judgement defaults for used techniques by category
    "inversion/GCD": ("mathematics", "T+Q", "ja", "schoon", 3, "O(n^2)", "O(n)", "O(n^2)", "ja", "Bennett", "nee", "laag", 4, 4),
    "scheduling/fusion": ("computer science", "T", "nee", "geen", 2, "-", "-", "-", "nee", "geen", "nee", "laag", 3, 3),
    "width/qubit management": ("computer science", "Q", "nee", "vuil", 2, "O(1)", "-", "-", "nee", "geen", "nee", "midden", 3, 3),
    "approximation": ("statistics", "T", "nee", "geen", 1, "-", "-", "-", "ja", "geen", "nee", "laag", 2, 3),
    "reduction": ("mathematics", "T", "ja", "schoon", 2, "O(n)", "O(n)", "O(n)", "nee", "Bennett", "ja", "laag", 3, 4),
    "validation/nonce": ("cryptography", "geen", "nee", "geen", 1, "-", "-", "-", "nee", "geen", "nee", "laag", 2, 3),
    "uncomputation": ("physics", "T", "nee", "geen", 1, "O(n)", "O(1)", "O(n)", "nee", "meting", "nee", "hoog", 3, 4),
    "encoding/transcript": ("information theory", "Q", "ja", "schoon", 2, "O(n)", "O(n)", "O(n)", "ja", "Bennett", "nee", "laag", 3, 3),
    "adder": ("computer science", "T", "ja", "schoon", 2, "O(n)", "O(n)", "O(n)", "nee", "meting", "nee", "laag", 2, 5),
    "multiplication": ("mathematics", "T", "ja", "schoon", 3, "O(n^1.58)", "O(n)", "O(n)", "nee", "Bennett", "nee", "laag", 3, 5),
    "classical/precompute": ("computer science", "T", "nee", "geen", 1, "-", "-", "-", "nee", "geen", "ja", "laag", 2, 4),
}


def used_rows():
    d = json.load(open(f"{SCR}/used_techniques.json"))
    rows = []
    for t in d:
        (field, goal, fixed, anc, coh, ts, qs, ds, dd, unc, tab, ph, cx, mat) = CAT_DEFAULTS[t["category"]]
        meas = str(t["uses_measurement"]).lower()
        midm = "ja" if meas.startswith(("yes", "partial", "classical")) else "nee"
        ff = "ja" if ("condition" in meas or "cz" in meas) else "nee"
        if midm == "ja" and unc == "geen":
            unc = "meting"
        inl = {"yes": "ja", "partial": "deels", "no": "nee"}[t["in_leader"]]
        note = f"{t['description']} | effect: {t['effect']} | {t.get('notes', '')}"
        rows.append([t["name"], t["category"], field, str(t.get("date", ""))[:4] or "2026",
                     f"ecdsa.fail {t.get('first_id', '')} ({t.get('solver', '')}, {str(t.get('date', ''))[:10]})",
                     goal, fixed, "-", anc, coh, "nee", midm, ff, "benaderd" if t["approximate"] else "exact",
                     "per shot (lambda)" if t["approximate"] else "0", ts, qs, ds, "ja", "nee", dd, unc, tab, ph, cx, mat,
                     0 if inl == "ja" else "-", 0 if inl == "ja" else "-", inl, "ja", note[:600]])
    return rows


def unused_rows():
    return [list(r) for r in UNUSED]


def textbook_rows():
    return [[f"[{r[0] + ' qubits' if r[0] != 'n' else 'n-bit'}] {r[1]}"] + list(r[2:31]) + [f"{r[31]} | in nieuw ontwerp: {r[32]}"] for r in TEXTBOOK]


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR
        cell.fill = HFILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_catalog(wb, title, prefix, rows, intro):
    ws = wb.create_sheet(title)
    ws["A1"] = intro
    ws["A1"].font = BOLD
    hr = 3
    for i, h in enumerate(COLS, 1):
        ws.cell(row=hr, column=i, value=h)
    for j, (name, _, _) in enumerate(FEATS):
        ws.cell(row=hr, column=FEAT0 + j, value=f"f:{name}")
    style_header(ws, hr, FEAT0 + NF - 1)
    for k, r in enumerate(rows):
        rr = hr + 1 + k
        ws.cell(row=rr, column=1, value=f"{prefix}{k + 1:02d}").font = BLACK
        for i, v in enumerate(r, 2):
            c = ws.cell(row=rr, column=i, value=v)
            c.font = BLUE
            c.alignment = Alignment(wrap_text=(i in (2, 6, 32)), vertical="top")
        for j, (_, col, tmpl) in enumerate(FEATS):
            ref = f"{get_column_letter(col)}{rr}"
            ws.cell(row=rr, column=FEAT0 + j, value="=" + tmpl.format(c=ref)).font = BLACK
    widths = {1: 7, 2: 42, 3: 16, 4: 16, 5: 7, 6: 34, 32: 70}
    for i in range(1, FEAT0 + NF):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 12)
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    return ws, hr + 1, hr + len(rows)


def main():
    wb = Workbook()
    leg = wb.active
    leg.title = "Legenda"
    used = used_rows()
    unused = unused_rows()
    text = textbook_rows()
    lines = [
        ("Technieken-analyse ecdsa.fail (secp256k1 point add), 26-10-2026", BOLD),
        ("Score = gemiddeld uitgevoerde Toffoli x piek-qubits. Leader 9ef3139: 741,645 x 1,234 = 915.2M.", BLACK),
        ("", BLACK),
        ("Bladen", BOLD),
        ("Gebruikt: 45 technieken uit alle oplossingen (154 inzendingsnotities gelezen + leader-broncode).", BLACK),
        ("Ongebruikt: 36 quantum-computing technieken die de leader niet gebruikt, incl. recent werk (2024-2026).", BLACK),
        ("Tekstboek: 35 methoden van 2-qubit-poorten tot 256-bit bouwstenen; kolom Opmerking zegt of ze in het nieuwe ontwerp zitten.", BLACK),
        ("Weging: extra eigenschappen + gewichtenmodel (gele cellen aanpasbaar); score per kandidaat is een formule.", BLACK),
        ("Gelijkheid: Gower-gelijkenis van elke kandidaat met elke gebruikte techniek (14 numerieke kenmerken, formules).", BLACK),
        ("Ontwerp: het nieuwe ontwerp zonder leader-constructies, opgebouwd uit primitieven (formules; blauwe cellen = invoer).", BLACK),
        ("Vergelijking: leader, Google (README), gepubliceerd, nieuwe varianten.", BLACK),
        ("Mijlpalen: alle stappen die >=1% van de beste score afhaalden.", BLACK),
        ("", BLACK),
        ("Kleuren", BOLD),
        ("Blauw = ingevoerde waarde (bron of oordeel); zwart = formule; groen = verwijzing naar ander blad; geel = gewicht om aan te passen.", BLACK),
        ("", BLACK),
        ("Status van de waarden", BOLD),
        ("Gemeten in deze repo: transcript lokaal 694 bits; Fermat-pebbling 961B; divsteps 591 stappen, max 535 op 1M invoeren; railbreedtes.", BLACK),
        ("Uit bronnen: Google-punten (README), gepubliceerde scores, Toffoli-formules van primitieven (citaten in kolom Bron).", BLACK),
        ("Oordeel (schatting): coherentie, complexiteit, volwassenheid, T/Q-effect op de leader, extra eigenschappen in Weging.", BLACK),
        ("Gebruikte technieken: kenmerken per categorie afgeleid (CAT_DEFAULTS in build_workbook.py) + vlaggen uit de notities.", BLACK),
        ("Het ontwerp is een kostenmodel, geen gebouwd circuit: het is niet door eval_circuit gevalideerd.", BLACK),
    ]
    for i, (t, f) in enumerate(lines, 1):
        leg.cell(row=i, column=1, value=t).font = f
    leg.column_dimensions["A"].width = 130

    ws_u, u0, u1 = write_catalog(wb, "Gebruikt", "G", used, "Technieken in alle ingediende oplossingen (inclusief leader)")
    ws_n, n0, n1 = write_catalog(wb, "Ongebruikt", "U", unused, "Technieken die de leader niet gebruikt (incl. recent academisch werk)")
    ws_t, t0, t1 = write_catalog(wb, "Tekstboek", "B", text, "Tekstboekmethoden van 2 qubits tot 256-bit bouwstenen")

    # ---- Weging -------------------------------------------------------------------------------------------------
    wg = wb.create_sheet("Weging")
    wg["A1"] = "Gewichtenmodel: gewogen geschiktheid van elke kandidaat (Ongebruikt + Tekstboek). Pas de gele gewichten aan."
    wg["A1"].font = BOLD
    weights = [("T-winst (T-effect, -10% = 1)", 3), ("Q-winst (Q-effect, -200 = 1)", 3), ("Exact", 2),
               ("Challenge-gateset", 3), ("Volwassenheid", 1), ("Lage complexiteit", 1), ("Nieuwheid (Gelijkheid)", 1),
               ("Toepasbaar op inversie", 2), ("Toepasbaar op payload", 1), ("Toepasbaar op breedte", 1),
               ("Combineerbaar met leader", 1), ("Bewijs / meting", 1)]
    wg["A3"] = "Kenmerk"
    wg["B3"] = "Gewicht"
    style_header(wg, 3, 2)
    for i, (k, w) in enumerate(weights):
        wg.cell(row=4 + i, column=1, value=k).font = BLACK
        c = wg.cell(row=4 + i, column=2, value=w)
        c.font = BLUE
        c.fill = YELLOW
    wsum_row = 4 + len(weights)
    wg.cell(row=wsum_row, column=1, value="Som gewichten").font = BOLD
    wg.cell(row=wsum_row, column=2, value=f"=SUM(B4:B{wsum_row - 1})").font = BLACK
    h0 = wsum_row + 2
    heads = ["ID", "Techniek", "Bron-blad", "f T-winst", "f Q-winst", "f exact", "f gateset", "f volwassen",
             "f eenvoud", "f nieuwheid", "Inversie (0-1)", "Payload (0-1)", "Breedte (0-1)", "Combineerbaar (0-1)",
             "Bewijs (0/0.5/1)", "Gewogen score", "Rang"]
    for i, h in enumerate(heads, 1):
        wg.cell(row=h0, column=i, value=h)
    style_header(wg, h0, len(heads))
    cands = [("Ongebruikt", n0, k) for k in range(n1 - n0 + 1)] + [("Tekstboek", t0, k) for k in range(t1 - t0 + 1)]
    measured = {"U01", "U07", "U09", "U10", "U31", "U32", "U33", "U35"}
    first = h0 + 1
    last = h0 + len(cands)
    for k, (sheet, r0, off) in enumerate(cands):
        rr = first + k
        src = r0 + off
        name_src = UNUSED[off][0] if sheet == "Ongebruikt" else text[off][0]
        cid = ("U" if sheet == "Ongebruikt" else "B") + f"{off + 1:02d}"
        cat = (UNUSED[off][1] if sheet == "Ongebruikt" else text[off][1]).lower()
        ref = lambda col: f"'{sheet}'!{get_column_letter(col)}{src}"
        wg.cell(row=rr, column=1, value=f"={ref(1)}").font = GREEN
        wg.cell(row=rr, column=2, value=f"={ref(2)}").font = GREEN
        wg.cell(row=rr, column=3, value=sheet).font = BLACK
        wg.cell(row=rr, column=4, value=f"=IF(ISNUMBER({ref(28)}),MAX(-1,MIN(1,-{ref(28)}/10)),0)").font = BLACK
        wg.cell(row=rr, column=5, value=f"=IF(ISNUMBER({ref(29)}),MAX(-1,MIN(1,-{ref(29)}/200)),0)").font = BLACK
        wg.cell(row=rr, column=6, value=f'=IF({ref(15)}="exact",1,0)').font = BLACK
        wg.cell(row=rr, column=7, value=f'=IF({ref(20)}="ja",1,0)').font = BLACK
        wg.cell(row=rr, column=8, value=f"={ref(27)}/5").font = BLACK
        wg.cell(row=rr, column=9, value=f"=1-{ref(26)}/5").font = BLACK
        wg.cell(row=rr, column=10, value=f"=INDEX(Gelijkheid!$C:$C,MATCH(A{rr},Gelijkheid!$A:$A,0))").font = GREEN
        inv = 1 if cat in ("inversion",) else 0.3 if cat in ("uncomputation", "encoding", "adder", "reduction") else 0
        pay = 1 if cat in ("reduction", "multiplication", "adder") else 0.3 if cat in ("inversion", "lookup") else 0
        wid = 1 if cat in ("width", "uncomputation", "encoding") else 0.3 if cat in ("adder", "comparison") else 0
        comb = 0 if cat in ("algorithm", "curve model", "validation") else 0.5
        if sheet == "Ongebruikt" and UNUSED[off][27] == "n.a.":
            comb = 0
        proof = 1 if cid in measured else 0.5 if sheet == "Tekstboek" else 0
        for j, v in enumerate([inv, pay, wid, comb, proof]):
            c = wg.cell(row=rr, column=11 + j, value=v)
            c.font = BLUE
        wg.cell(row=rr, column=16, value=(
            f"=($B$4*D{rr}+$B$5*E{rr}+$B$6*F{rr}+$B$7*G{rr}+$B$8*H{rr}+$B$9*I{rr}+$B$10*J{rr}"
            f"+$B$11*K{rr}+$B$12*L{rr}+$B$13*M{rr}+$B$14*N{rr}+$B$15*O{rr})/$B${wsum_row}")).font = BLACK
        wg.cell(row=rr, column=17, value=f"=RANK(P{rr},$P${first}:$P${last})").font = BLACK
    for i, wdt in enumerate([7, 46, 12] + [10] * 14, 1):
        wg.column_dimensions[get_column_letter(i)].width = wdt
    wg.freeze_panes = wg.cell(row=first, column=3)
    wg.cell(row=last + 2, column=1, value=(
        "Extra eigenschappen (Inversie/Payload/Breedte/Combineerbaar/Bewijs) zijn oordelen per categorie; "
        "Bewijs = 1 als in deze repo gemeten, 0.5 voor tekstboek-primitieven met bekende formule.")).font = BLACK

    # ---- Gelijkheid ---------------------------------------------------------------------------------------------
    gl = wb.create_sheet("Gelijkheid")
    gl["A1"] = ("Gower-gelijkenis = 1 - gemiddelde |verschil| over 14 kenmerken (f:-kolommen van de catalogi). "
                "Nieuwheid = 1 - max gelijkenis met de gebruikte set.")
    gl["A1"].font = BOLD
    hr = 3
    gl.cell(row=hr, column=1, value="ID")
    gl.cell(row=hr, column=2, value="Techniek")
    gl.cell(row=hr, column=3, value="Nieuwheid")
    gl.cell(row=hr, column=4, value="Max gelijkenis")
    gl.cell(row=hr, column=5, value="Meest gelijkende gebruikte techniek")
    m0 = 6
    nused = u1 - u0 + 1
    for j in range(nused):
        gl.cell(row=hr, column=m0 + j, value=f"=Gebruikt!A{u0 + j}").font = HDR
    style_header(gl, hr, m0 + nused - 1)
    fa = get_column_letter(FEAT0)
    fb = get_column_letter(FEAT0 + NF - 1)
    for k, (sheet, r0, off) in enumerate(cands):
        rr = hr + 1 + k
        src = r0 + off
        gl.cell(row=rr, column=1, value=f"='{sheet}'!A{src}").font = GREEN
        gl.cell(row=rr, column=2, value=f"='{sheet}'!B{src}").font = GREEN
        for j in range(nused):
            ur = u0 + j
            gl.cell(row=rr, column=m0 + j, value=(
                f"=1-SUMPRODUCT(ABS('{sheet}'!${fa}${src}:${fb}${src}-Gebruikt!${fa}${ur}:${fb}${ur}))/{NF}")).font = BLACK
        rng = f"{get_column_letter(m0)}{rr}:{get_column_letter(m0 + nused - 1)}{rr}"
        gl.cell(row=rr, column=4, value=f"=MAX({rng})").font = BLACK
        gl.cell(row=rr, column=3, value=f"=1-D{rr}").font = BLACK
        gl.cell(row=rr, column=5, value=f"=INDEX(Gebruikt!$B${u0}:$B${u1},MATCH(D{rr},{rng},0))").font = BLACK
    gl.column_dimensions["B"].width = 46
    gl.column_dimensions["E"].width = 46
    gl.freeze_panes = gl.cell(row=hr + 1, column=3)

    # ---- Ontwerp ------------------------------------------------------------------------------------------------
    on = wb.create_sheet("Ontwerp")
    on["A1"] = "Nieuw ontwerp 'DivPay': half-delta divsteps + mod-p payload, uit tekstboek-primitieven (kostenmodel)"
    on["A1"].font = BOLD
    inputs = [
        ("n (bits)", 256, "secp256k1"),
        ("Stappen T6", 535, "steps_tail.py: max 535 op 1,000,000 invoeren (P(>534)=4e-6)"),
        ("Stappen strikt", 591, "bewezen grens (sipa safegcd hull)"),
        ("Som railbreedte-envelop (535)", 72315, "divsteps_sim.py envelop (1,500 invoeren, +4 bits marge)"),
        ("Som ceil(w/16) (535)", 4775, "idem, voor de opgeknipte optellers"),
        ("Som railbreedte-envelop (591)", 72651, "idem"),
        ("Railbreedte laatste stap", 6, "idem"),
        ("Blokgrootte opgeknipte opteller k", 16, "oordeel"),
        ("Vergelijkvenster w (bits)", 40, "foutkans ~2^-40 per vergelijking"),
        ("Delta- en letterlogica per stap (T)", 42, "30 delta + 12 trit-pakking (schatting)"),
        ("Trit-pakking bits per stap", 27 / 17, "17 trits in 27 bits"),
        ("Vergelijker-ancilla (Khattar-Gidney)", 4, "log* n"),
    ]
    on["A3"] = "Invoer"
    on["B3"] = "Waarde"
    on["C3"] = "Bron / aanname"
    style_header(on, 3, 3)
    for i, (k, v, s) in enumerate(inputs):
        on.cell(row=4 + i, column=1, value=k).font = BLACK
        on.cell(row=4 + i, column=2, value=v).font = BLUE
        on.cell(row=4 + i, column=3, value=s).font = BLACK
    # named refs
    n, st, sts, sw, sc, sw591, wend, kk, win, dl, tp, ca = [f"$B${4 + i}" for i in range(len(inputs))]
    r = 4 + len(inputs) + 1
    on.cell(row=r, column=1, value="Component (Toffoli)").font = HDR
    on.cell(row=r, column=2, value="T6 (generiek, 2.95B)").font = HDR
    on.cell(row=r, column=3, value="Strikt (geen enkele leader-techniek)").font = HDR
    for c in range(1, 4):
        on.cell(row=r, column=c).fill = HFILL
    comps = [
        ("Opteller n bits (T)", f"={n}-1+2*ROUNDUP({n}/{kk},0)", f"=2*{n}-1"),
        ("Rails per pass", f"=2*{sw}-{st}+2*{sc}+{st}*{dl}", f"=3*{sw591}-{sts}+{sts}*{dl}"),
        ("Payload per stap: swap", f"={n}", f"={n}"),
        ("Payload per stap: add/sub", "=B{a}", "=C{a}"),
        ("Payload per stap: mod-p correctie", f"=4*{win}+{n}", f"=5*{n}"),
        ("Payload per stap: halveren", f"={n}+{win}", f"=2*{n}"),
        ("Payload per walk", "=B{st_}*(B{p1}+B{p2}+B{p3}+B{p4})", "=C{st_}*(C{p1}+C{p2}+C{p3}+C{p4})"),
        ("Walk (2 railpassen + payload)", "=2*B{rp}+B{pw}", "=2*C{rp}+C{pw}"),
        ("Kwadraat (schoolboek n^2/2 + reductie)", f"={n}^2/2+18*{n}+B{{a}}", f"={n}^2/2+{n}^2+C{{a}}"),
        ("Coordinaten (4 klassieke aftrekkingen)", f"=4*{n}", f"=4*{n}"),
        ("Totaal Toffoli", "=2*B{wk}+B{sq}+B{co}", "=2*C{wk}+C{sq}+C{co}"),
        ("Transcript (qubits)", f"=ROUNDUP({st}*{tp},0)", f"=ROUNDUP({sts}*{tp},0)"),
        ("Payload u,v (qubits)", f"=2*{n}", f"=2*{n}"),
        ("Rails einde + delta + misc (qubits)", f"=2*{wend}+11+8", f"=2*{wend}+11+8"),
        ("Opteller-ancilla (qubits)", f"={kk}+1", "=1"),
        ("Vergelijker-ancilla (qubits)", f"={ca}", "=4"),
        ("Piek-qubits", "=SUM(B{q1}:B{q5})", "=SUM(C{q1}:C{q5})"),
        ("Score (T x Q)", "=B{tt}*B{qq}", "=C{tt}*C{qq}"),
    ]
    base = r + 1
    idx = {name: base + i for i, (name, _, _) in enumerate(comps)}
    keys = dict(a=idx["Opteller n bits (T)"], rp=idx["Rails per pass"], p1=idx["Payload per stap: swap"],
                p2=idx["Payload per stap: add/sub"], p3=idx["Payload per stap: mod-p correctie"],
                p4=idx["Payload per stap: halveren"], pw=idx["Payload per walk"], wk=idx["Walk (2 railpassen + payload)"],
                sq=idx["Kwadraat (schoolboek n^2/2 + reductie)"], co=idx["Coordinaten (4 klassieke aftrekkingen)"],
                tt=idx["Totaal Toffoli"], q1=idx["Transcript (qubits)"], q5=idx["Vergelijker-ancilla (qubits)"],
                qq=idx["Piek-qubits"])
    for i, (name, f6, fs) in enumerate(comps):
        rr = base + i
        on.cell(row=rr, column=1, value=name).font = BOLD if name in ("Totaal Toffoli", "Piek-qubits", "Score (T x Q)") else BLACK
        b = f6.replace("{{a}}", str(keys["a"]))
        c_ = fs.replace("{{a}}", str(keys["a"]))
        b = b.replace("{st_}", "").replace("B{st_}", "")
        b = b.format(st_="", **keys) if "{" in b else b
        c_ = c_.format(st_="", **keys) if "{" in c_ else c_
        on.cell(row=rr, column=2, value=b).font = BLACK
        on.cell(row=rr, column=3, value=c_).font = BLACK
    # payload per walk uses the step count, not a row: patch those two cells
    pw = keys["pw"]
    on.cell(row=pw, column=2, value=f"={st}*(B{keys['p1']}+B{keys['p2']}+B{keys['p3']}+B{keys['p4']})")
    on.cell(row=pw, column=3, value=f"={sts}*(C{keys['p1']}+C{keys['p2']}+C{keys['p3']}+C{keys['p4']})")
    for c in (2, 3):
        for rr in range(base, base + len(comps)):
            on.cell(row=rr, column=c).number_format = "#,##0"
    on.column_dimensions["A"].width = 44
    on.column_dimensions["B"].width = 24
    on.column_dimensions["C"].width = 60
    on.cell(row=base + len(comps) + 1, column=1, value=(
        "T6 gebruikt generieke primitieven die de leader óók gebruikt (venster-vergelijkers #G17, opgeknipte optellers #G19, "
        "Solinas #G24, meting-uncompute #G14). Strikt vermijdt alle technieken met 'In leader' = ja/deels.")).font = BLACK
    score_t6 = f"Ontwerp!B{keys['qq'] + 1}"
    score_st = f"Ontwerp!C{keys['qq'] + 1}"
    t6_t, t6_q = f"Ontwerp!B{keys['tt']}", f"Ontwerp!B{keys['qq']}"
    st_t, st_q = f"Ontwerp!C{keys['tt']}", f"Ontwerp!C{keys['qq']}"

    # ---- Vergelijking -------------------------------------------------------------------------------------------
    vg = wb.create_sheet("Vergelijking")
    vg["A1"] = "Vergelijking: score = gemiddeld uitgevoerde Toffoli x piek-qubits"
    vg["A1"].font = BOLD
    hd = ["Oplossing", "Toffoli", "Qubits", "Score", "x leader", "x Google low-gate", "Exact?", "Bron"]
    for i, h in enumerate(hd, 1):
        vg.cell(row=3, column=i, value=h)
    style_header(vg, 3, len(hd))
    rows = [
        ("Leader 9ef3139 (benaderd, gegrind)", 741645, 1234, "nee (lambda~16 + nonce)", "ecdsafail submissions --all"),
        ("Leader-afgeleid exact (schatting)", 784645, 1243, "ja (lambda~0.2)", "docs/exactness-experiments.md (+43k T, ~+9 Q)"),
        ("Google low-gate", 2100000, 1425, "onbekend (privé)", "challenge README"),
        ("Google low-qubit", 2700000, 1175, "onbekend (privé)", "challenge README"),
        ("Long et al. 2026", 1356000, 1419, "nee (tegenvoorbeelden)", "arXiv 2609.28882"),
        ("Theta 893-qubit", 2218127, 893, "nee (111/360,960 fout)", "postquantum.com"),
        ("arXiv 2606.02235 space-opt", 2390000, 1192, "nee (benaderd)", "arXiv 2606.02235 (2^21.19)"),
        ("Challenge-startcircuit", 3942753, 2715, "ja", "challenge README"),
    ]
    for i, (nm, t, q, ex, src) in enumerate(rows):
        rr = 4 + i
        vg.cell(row=rr, column=1, value=nm).font = BLACK
        vg.cell(row=rr, column=2, value=t).font = BLUE
        vg.cell(row=rr, column=3, value=q).font = BLUE
        vg.cell(row=rr, column=7, value=ex).font = BLACK
        vg.cell(row=rr, column=8, value=src).font = BLACK
    rr = 4 + len(rows)
    for nm, tref, qref, ex in [("Nieuw: DivPay T6 (generieke primitieven)", t6_t, t6_q, "ja (lambda~0.02)"),
                               ("Nieuw: DivPay strikt (geen leader-techniek)", st_t, st_q, "ja (bewezen 591 stappen)")]:
        vg.cell(row=rr, column=1, value=nm).font = BOLD
        vg.cell(row=rr, column=2, value=f"={tref}").font = GREEN
        vg.cell(row=rr, column=3, value=f"={qref}").font = GREEN
        vg.cell(row=rr, column=7, value=ex).font = BLACK
        vg.cell(row=rr, column=8, value="Ontwerp-blad (kostenmodel, niet gebouwd)").font = BLACK
        rr += 1
    for r_ in range(4, rr):
        vg.cell(row=r_, column=4, value=f"=B{r_}*C{r_}").font = BLACK
        vg.cell(row=r_, column=5, value=f"=D{r_}/$D$4").font = BLACK
        vg.cell(row=r_, column=6, value=f"=D{r_}/$D$6").font = BLACK
        vg.cell(row=r_, column=2).number_format = "#,##0"
        vg.cell(row=r_, column=4).number_format = "#,##0"
        vg.cell(row=r_, column=5).number_format = "0.00x"
        vg.cell(row=r_, column=6).number_format = "0.000x"
    vg.cell(row=rr + 1, column=1, value="Verslaat T6 Google low-gate?").font = BOLD
    vg.cell(row=rr + 1, column=2, value=f'=IF(D{rr - 2}<D6,"ja, met "&ROUND((1-D{rr - 2}/D6)*100,1)&" procent","nee")').font = BLACK
    vg.cell(row=rr + 2, column=1, value="Verslaat strikt Google low-gate?").font = BOLD
    vg.cell(row=rr + 2, column=2, value=f'=IF(D{rr - 1}<D6,"ja","nee, "&ROUND(D{rr - 1}/D6,2)&"x zo hoog")').font = BLACK
    vg.cell(row=4, column=2).comment = Comment("Laatste promoted inzending 26-10-2026", "Claude")
    for i, wdt in enumerate([44, 14, 9, 18, 10, 16, 26, 48], 1):
        vg.column_dimensions[get_column_letter(i)].width = wdt

    # ---- Mijlpalen ----------------------------------------------------------------------------------------------
    ms = json.load(open(f"{SCR}/milestones.json"))
    mp = wb.create_sheet("Mijlpalen")
    mp["A1"] = f"Stappen >=1% t.o.v. de lopende beste score. Start {ms['start']['score']:,} -> {ms['end']['score']:,}"
    mp["A1"].font = BOLD
    drops = ms["drops_ge_1pct"]
    keys_m = list(drops[0].keys())
    for i, k in enumerate(keys_m, 1):
        mp.cell(row=3, column=i, value=k)
    style_header(mp, 3, len(keys_m))
    for j, dd in enumerate(drops):
        for i, k in enumerate(keys_m, 1):
            v = dd[k]
            mp.cell(row=4 + j, column=i, value=v if isinstance(v, (int, float, str)) or v is None else json.dumps(v)).font = BLUE
    for i in range(1, len(keys_m) + 1):
        mp.column_dimensions[get_column_letter(i)].width = 16
    mp.column_dimensions[get_column_letter(len(keys_m))].width = 70

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.font is None or c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None)
    wb.save(OUT)
    print("saved", OUT, "used", len(used), "unused", len(unused), "textbook", len(text))


if __name__ == "__main__":
    main()
